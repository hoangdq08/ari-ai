from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader


@dataclass
class IngestedDocument:
    source_id: str
    raw_text: str
    metadata: dict[str, Any]
    is_duplicate: bool = False
    duplicate_of: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class DataIngestor:
    def __init__(self, raw_dir: str | Path):
        self.raw_dir = Path(raw_dir)
        self.raw_dir.mkdir(parents=True, exist_ok=True)

    def ingest_file(
        self,
        file_bytes: bytes,
        file_name: str,
        source_type: str,
        title: str | None,
        reliability_level: str,
    ) -> IngestedDocument:
        suffix = Path(file_name).suffix.lower()
        if suffix == ".pdf":
            raw_text = _extract_pdf(file_bytes)
        elif suffix == ".docx":
            raw_text = _extract_docx(file_bytes)
        elif suffix in {".html", ".htm"}:
            raw_text = _extract_html(file_bytes.decode("utf-8", errors="ignore"))
        else:
            raw_text = file_bytes.decode("utf-8", errors="ignore")

        content_hash = _content_hash(raw_text)
        duplicate = self._find_duplicate(content_hash)
        if duplicate:
            duplicate_metadata, duplicate_text = duplicate
            return IngestedDocument(
                source_id=duplicate_metadata["source_id"],
                raw_text=duplicate_text,
                metadata=duplicate_metadata,
                is_duplicate=True,
                duplicate_of=duplicate_metadata["source_id"],
            )

        source_id = _source_id(f"file:{content_hash}")
        metadata = {
            "source_id": source_id,
            "source_type": source_type,
            "title": title or Path(file_name).stem,
            "file_name": file_name,
            "url": None,
            "page": None,
            "crawled_at": datetime.now(timezone.utc).isoformat(),
            "reliability_level": reliability_level,
            "content_hash": content_hash,
            "dedupe_key": f"sha256:{content_hash}",
        }
        self._save_raw(source_id, raw_text, metadata)
        return IngestedDocument(source_id=source_id, raw_text=raw_text, metadata=metadata)

    def ingest_url(self, url: str, title: str | None, reliability_level: str) -> IngestedDocument:
        response = requests.get(url, timeout=20, headers={"User-Agent": "NongTriAI/0.1"})
        response.raise_for_status()
        parsed = urlparse(url)
        content_type = response.headers.get("content-type", "").lower()
        suffix = Path(parsed.path).suffix.lower()

        if "pdf" in content_type or suffix == ".pdf":
            source_type = "pdf"
            raw_text = _extract_pdf(response.content)
            file_name = Path(parsed.path).name or f"{_source_id(url)}.pdf"
        elif "word" in content_type or suffix == ".docx":
            source_type = "docx"
            raw_text = _extract_docx(response.content)
            file_name = Path(parsed.path).name or f"{_source_id(url)}.docx"
        elif "text/plain" in content_type or suffix == ".txt":
            source_type = "txt"
            raw_text = response.text
            file_name = Path(parsed.path).name or f"{_source_id(url)}.txt"
        else:
            source_type = "web"
            raw_text = _extract_html(response.text)
            file_name = None

        content_hash = _content_hash(raw_text)
        source_id = _source_id(url)
        duplicate = self._find_duplicate(content_hash, allow_source_id=source_id)
        if duplicate:
            duplicate_metadata, duplicate_text = duplicate
            return IngestedDocument(
                source_id=duplicate_metadata["source_id"],
                raw_text=duplicate_text,
                metadata=duplicate_metadata,
                is_duplicate=True,
                duplicate_of=duplicate_metadata["source_id"],
            )

        existing_path = self.raw_dir / f"{source_id}.json"
        if existing_path.exists():
            existing_payload = _read_raw(existing_path)
            existing_metadata = existing_payload.get("metadata") or {}
            if existing_metadata.get("content_hash") == content_hash:
                return IngestedDocument(
                    source_id=source_id,
                    raw_text=existing_payload.get("raw_text") or raw_text,
                    metadata=existing_metadata,
                    is_duplicate=True,
                    duplicate_of=source_id,
                )

        metadata = {
            "source_id": source_id,
            "source_type": source_type,
            "title": title or parsed.netloc,
            "file_name": file_name,
            "url": url,
            "page": None,
            "crawled_at": datetime.now(timezone.utc).isoformat(),
            "reliability_level": reliability_level,
            "content_type": content_type,
            "content_hash": content_hash,
            "dedupe_key": f"url:{url}|sha256:{content_hash}",
        }
        self._save_raw(source_id, raw_text, metadata)
        return IngestedDocument(source_id=source_id, raw_text=raw_text, metadata=metadata)

    def _save_raw(self, source_id: str, raw_text: str, metadata: dict[str, Any]) -> None:
        with (self.raw_dir / f"{source_id}.json").open("w", encoding="utf-8") as handle:
            json.dump({"metadata": metadata, "raw_text": raw_text}, handle, ensure_ascii=False, indent=2)

    def _find_duplicate(self, content_hash: str, allow_source_id: str | None = None) -> tuple[dict[str, Any], str] | None:
        for path in sorted(self.raw_dir.glob("*.json")):
            payload = _read_raw(path)
            metadata = payload.get("metadata") or {}
            source_id = metadata.get("source_id") or path.stem
            if allow_source_id and source_id == allow_source_id:
                continue
            existing_hash = metadata.get("content_hash") or _content_hash(payload.get("raw_text") or "")
            if existing_hash == content_hash:
                metadata = {**metadata, "source_id": source_id}
                return metadata, payload.get("raw_text") or ""
        return None


def _extract_pdf(file_bytes: bytes) -> str:
    temp_path = Path("/private/tmp") / f"nong_tri_{_source_id(str(len(file_bytes)))}.pdf"
    temp_path.write_bytes(file_bytes)
    reader = PdfReader(str(temp_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(file_bytes: bytes) -> str:
    temp_path = Path("/private/tmp") / f"nong_tri_{_source_id(str(len(file_bytes)))}.docx"
    temp_path.write_bytes(file_bytes)
    document = Document(str(temp_path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_html(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript", "svg"]):
        tag.decompose()

    content_candidates = soup.select(
        "article, main, [role='main'], .entry-content, .post-content, .article-content, "
        ".single-content, .content-area, .td-post-content, .blog-detail, .blog-content"
    )
    if content_candidates:
        content = max(content_candidates, key=lambda node: len(node.get_text(" ", strip=True)))
        return content.get_text("\n")

    return soup.get_text("\n")


def _source_id(seed: str) -> str:
    return hashlib.sha1(seed.encode("utf-8")).hexdigest()[:16]


def _content_hash(text: str) -> str:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    normalized = " ".join(normalized.split()).casefold()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _read_raw(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        return payload if isinstance(payload, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}
